from __future__ import annotations

from io import BytesIO
from pathlib import Path
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook
from werkzeug.datastructures import FileStorage

from mcorsi.extensions import db
from mcorsi.models import Role, StoredFile, User
from mcorsi.services.storage import StorageError, path_for, save_upload


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _actor() -> User:
    actor = User(email="storage@example.it", profile_completed=True)
    actor.roles.append(Role.query.filter_by(name="operator").one())
    db.session.add(actor)
    db.session.flush()
    return actor


def _docx_bytes() -> bytes:
    output = BytesIO()
    document = Document()
    document.add_paragraph("Documento valido")
    document.save(output)
    return output.getvalue()


def _xlsx_bytes() -> bytes:
    output = BytesIO()
    workbook = Workbook()
    workbook.active.append(["dato"])
    workbook.save(output)
    return output.getvalue()


def test_client_mime_is_ignored_and_canonical_mime_is_stored(app):
    with app.app_context():
        stored = save_upload(
            FileStorage(
                stream=BytesIO(_docx_bytes()),
                filename="documento.docx",
                content_type="application/octet-stream",
            ),
            actor=_actor(),
            category="course-documents",
            allowed_mime_types={DOCX_MIME},
        )
        db.session.commit()

        assert stored.mime_type == DOCX_MIME
        assert path_for(stored).is_file()


@pytest.mark.parametrize(
    ("filename", "content"),
    [
        ("programma.pdf", b"MZ\x90\x00eseguibile"),
        ("immagine.png", b"\x89PNG\r\n\x1a\ncontenuto non valido"),
        ("testo.txt", b"testo\x00binario"),
    ],
)
def test_fake_signatures_are_rejected_and_removed(app, filename, content):
    with app.app_context():
        with pytest.raises(StorageError):
            save_upload(
                FileStorage(
                    stream=BytesIO(content),
                    filename=filename,
                    content_type="application/octet-stream",
                ),
                actor=_actor(),
                category="course-documents",
            )

        assert StoredFile.query.count() == 0
        storage = Path(app.config["PRIVATE_STORAGE_PATH"])
        assert not any(path.is_file() for path in storage.rglob("*"))


def test_office_container_must_match_its_extension(app):
    with app.app_context():
        with pytest.raises(StorageError, match="estensione"):
            save_upload(
                FileStorage(
                    stream=BytesIO(_xlsx_bytes()),
                    filename="falso.docx",
                    content_type=DOCX_MIME,
                ),
                actor=_actor(),
                category="course-documents",
            )


def test_macro_payload_is_rejected_even_with_docx_extension(app):
    output = BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
        archive.writestr(
            "word/document.xml",
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" />',
        )
        archive.writestr("word/vbaProject.bin", b"macro")

    with app.app_context():
        with pytest.raises(StorageError, match="macro"):
            save_upload(
                FileStorage(
                    stream=BytesIO(output.getvalue()),
                    filename="macro.docx",
                    content_type=DOCX_MIME,
                ),
                actor=_actor(),
                category="templates",
            )
