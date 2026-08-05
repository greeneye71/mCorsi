from __future__ import annotations

from datetime import date, datetime, time, timezone
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from PIL import Image
from flask import current_app
from reportlab.pdfgen import canvas
from werkzeug.datastructures import FileStorage

from mcorsi.extensions import db
from mcorsi.models import (
    Certificate,
    CertificateTemplate,
    Enrollment,
    ParticipantProfile,
    Questionnaire,
    QuestionnaireAttempt,
    Role,
    User,
)
from mcorsi.services.certificates import (
    CertificateError,
    generate_certificate,
    inspect_template,
    readiness,
    validate_signature_image,
)
from mcorsi.services.courses import create_course
from mcorsi.services.storage import StorageError, path_for, save_upload


def _user(role: str, email: str, *, first_name="", last_name="") -> User:
    user = User(
        email=email,
        first_name=first_name,
        last_name=last_name,
        profile_completed=True,
    )
    if role in {"admin", "operator"}:
        user.set_password("PasswordMoltoSicura!")
    user.roles.append(Role.query.filter_by(name=role).one())
    if role == "admin":
        user.roles.append(Role.query.filter_by(name="operator").one())
    db.session.add(user)
    db.session.flush()
    return user


def _pdf_bytes(text="attestato") -> bytes:
    output = BytesIO()
    pdf = canvas.Canvas(output)
    pdf.drawString(72, 760, text)
    pdf.showPage()
    pdf.save()
    return output.getvalue()


def test_default_docx_has_required_placeholders_and_geometry():
    path = Path("mcorsi/assets/default_certificate.docx").resolve()
    variables = inspect_template(path)
    assert "participant_full_name" in variables
    assert "course_title" in variables
    assert "signature_image" in variables
    document = Document(path)
    section = document.sections[0]
    assert round(section.page_width.mm) == 297
    assert round(section.page_height.mm) == 210
    assert len(document.tables) == 1


def test_storage_rejects_macro_documents_and_tiny_signatures(app):
    with app.app_context():
        admin = _user("admin", "security@example.it")
        with pytest.raises(StorageError):
            save_upload(
                FileStorage(
                    stream=BytesIO(b"macro"),
                    filename="documento.docm",
                    content_type="application/vnd.ms-word.document.macroEnabled.12",
                ),
                actor=admin,
                category="course-documents",
            )
        image = Image.new("RGB", (10, 10), "white")
        path = Path(app.config["PRIVATE_STORAGE_PATH"]) / "tiny.png"
        path.parent.mkdir(parents=True, exist_ok=True)
        image.save(path)
        with pytest.raises(CertificateError, match="troppo piccola"):
            validate_signature_image(path)


def test_generates_immutable_pdf_when_requirements_are_met(app, monkeypatch):
    with app.app_context():
        admin = _user("admin", "admin@example.it", first_name="Ada")
        participant = _user(
            "participant", "mario@example.it", first_name="Mario", last_name="Rossi"
        )
        participant.participant_profile = ParticipantProfile(
            birth_place="Roma", birth_date=date(1980, 5, 10), tax_code="RSSMRA80E10H501Z"
        )
        template_path = Path("mcorsi/assets/default_certificate.docx").resolve()
        stored_template = save_upload(
            FileStorage(
                stream=BytesIO(template_path.read_bytes()),
                filename="modello.docx",
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
            actor=admin,
            category="templates",
        )
        template = CertificateTemplate(
            name="Standard",
            stored_file=stored_template,
            placeholders=inspect_template(path_for(stored_template)),
            uploaded_by_user_id=admin.id,
        )
        course = create_course(
            actor=admin,
            data={
                "title": "Radioprotezione",
                "description": "",
                "status": "completed",
                "referent_user_id": admin.id,
                "session_date": date(2026, 3, 15),
                "start_time": time(9),
                "end_time": time(13),
                "delivery_mode": "online",
                "meeting_url": "",
                "certificate_validity_months": 60,
            },
        )
        course.certificate_template = template
        questionnaire = Questionnaire(
            course=course,
            title="Valutazione",
            passing_percentage=70,
            max_attempts=3,
            sort_order=1,
            is_published=True,
        )
        enrollment = Enrollment(
            course=course, participant=participant, attendance_status="attended"
        )
        db.session.add_all([template, questionnaire, enrollment])
        db.session.flush()
        db.session.add(
            QuestionnaireAttempt(
                questionnaire=questionnaire,
                participant=participant,
                attempt_number=1,
                submitted_at=datetime.now(timezone.utc),
                score=10,
                maximum_score=10,
                passing_percentage_snapshot=70,
                passed=True,
            )
        )
        db.session.commit()

        def fake_convert(_docx_path, output_dir):
            output_dir.mkdir(parents=True, exist_ok=True)
            result = output_dir / "attestato.pdf"
            result.write_bytes(_pdf_bytes())
            return result

        monkeypatch.setattr("mcorsi.services.certificates.convert_docx_to_pdf", fake_convert)
        assert readiness(enrollment) == (True, [])
        certificate = generate_certificate(enrollment, actor=admin)
        db.session.commit()
        assert certificate.certificate_number.startswith("MC-2026-")
        assert certificate.expires_at == date(2031, 3, 15)
        assert certificate.data_snapshot["participant_full_name"] == "Mario Rossi"
        assert path_for(certificate.pdf_file).read_bytes().startswith(b"%PDF")


def test_participant_upload_is_pending_and_private(app, client):
    with app.app_context():
        participant = _user(
            "participant", "anna@example.it", first_name="Anna", last_name="Bianchi"
        )
        other = _user("participant", "altro@example.it")
        db.session.commit()
        participant_id, other_id = participant.id, other.id
    with client.session_transaction() as browser_session:
        browser_session["_user_id"] = participant_id
        browser_session["_fresh"] = True
    response = client.post(
        "/documents/participant/upload",
        data={
            "title": "Primo soccorso",
            "course_date": "2024-02-20",
            "expires_at": "2027-02-20",
            "file": (BytesIO(_pdf_bytes("storico")), "attestato.pdf"),
        },
        content_type="multipart/form-data",
    )
    assert response.status_code == 302
    with app.app_context():
        certificate = Certificate.query.one()
        assert certificate.verification_status == "pending"
        certificate_id = certificate.id
    assert client.get(f"/documents/certificates/{certificate_id}/download").status_code == 200
    with client.session_transaction() as browser_session:
        browser_session["_user_id"] = other_id
        browser_session["_fresh"] = True
    assert client.get(f"/documents/certificates/{certificate_id}/download").status_code == 403
